"""端到端验证脚本：构造真实 EPUB/DOCX，经引擎转换到全部输出格式。"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


class PrintListener:
    def on_task_started(self, path):
        print(f"  开始: {Path(path).name}")

    def on_task_progress(self, path, ratio, message):
        pass

    def on_task_finished(self, path, ok, message, word_count):
        mark = "✓" if ok else "✗"
        wc = f"，字数 {word_count}" if word_count else ""
        print(f"  {mark} {message}{wc}")

    def on_batch_finished(self, success_count, total):
        print(f"批次完成: {success_count}/{total}")
        assert success_count == total, "存在失败任务"


def make_sample_epub(path):
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("e2e-001")
    book.set_title("端到端测试书")
    book.set_language("zh")
    book.add_author("测试作者")
    spine = ["nav"]
    for i in range(1, 4):
        c = epub.EpubHtml(title=f"第{i}章", file_name=f"ch{i}.xhtml", lang="zh")
        c.set_content(f"<h1>第{i}章</h1><p>这是第{i}章的内容，包含中文与 English 混排。</p>")
        book.add_item(c)
        spine.append(c)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = spine
    epub.write_epub(str(path), book)


def make_sample_docx(path):
    from docx import Document

    doc = Document()
    doc.core_properties.title = "DOCX 测试书"
    doc.add_heading("第一章", level=1)
    doc.add_paragraph("DOCX 第一章正文。")
    doc.add_heading("第二章", level=1)
    doc.add_paragraph("DOCX 第二章正文。")
    doc.save(str(path))


def main():
    import tempfile

    from core.engine import ConversionEngine, ConversionOptions

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        inputs = tmp / "inputs"
        inputs.mkdir()
        output = tmp / "output"

        epub_path = inputs / "sample.epub"
        make_sample_epub(epub_path)
        docx_path = inputs / "sample.docx"
        make_sample_docx(docx_path)

        for fmt in ("txt", "md", "html", "epub"):
            print(f"\n== 输出格式: {fmt} ==")
            out = output / fmt
            engine = ConversionEngine(PrintListener(), max_workers=2)
            engine.submit([str(epub_path), str(docx_path)], ConversionOptions(str(out), fmt))
            assert engine.wait(timeout=120), "转换超时"
            engine.shutdown(wait=True)
            produced = sorted(p.name for p in out.iterdir())
            print(f"  产物: {produced}")
            assert len(produced) == 2, f"预期 2 个产物，实际 {produced}"

        # 验证 EPUB→TXT 的章节顺序与元信息（同名 stem 并发转换时序不定，按内容定位产物）
        txt_files = list((output / "txt").glob("*.txt"))
        epub_txt = next(p for p in txt_files if "端到端测试书" in p.read_text(encoding="utf-8"))
        txt = epub_txt.read_text(encoding="utf-8")
        assert "书名：端到端测试书" in txt
        assert "作者：测试作者" in txt
        assert txt.index("第1章") < txt.index("第2章") < txt.index("第3章"), "章节顺序错误"

        # 验证 EPUB→EPUB 可被再次读取
        from core.extractors.registry import get_extractor

        # 重新读取 epub 产物中内容匹配的那个
        target = None
        for p in (output / "epub").glob("*.epub"):
            result = get_extractor(str(p)).extract(str(p))
            if len(result.chapters) == 3 and "第3章" in result.chapters[2].text:
                target = p
                break
        assert target is not None, "未找到有效的 EPUB 产物"

        print("\n✓ 端到端验证全部通过")


if __name__ == "__main__":
    main()
