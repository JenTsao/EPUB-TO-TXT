#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DOCX 提取器（python-docx）：按 Heading 样式切分章节。"""

from core.extractors.base import BaseExtractor, ConversionError, TaskControl
from core.models import BookMetadata, Chapter, ExtractionResult
from core.text import clean_text


def _style_name(paragraph) -> str:
    try:
        return (paragraph.style.name or "").lower() if paragraph.style else ""
    except Exception:
        return ""


class DocxExtractor(BaseExtractor):
    def _open(self, path: str):
        try:
            from docx import Document
        except ImportError as e:
            raise ConversionError("未安装 python-docx 库，无法处理 DOCX 文件（pip install python-docx）") from e
        try:
            return Document(path)
        except Exception as e:
            raise ConversionError(f"DOCX 打开失败：{e}") from e

    def metadata(self, path: str) -> BookMetadata:
        meta = BookMetadata.from_path(path)
        try:
            document = self._open(path)
            props = document.core_properties
            meta.title = props.title or None
            meta.author = props.author or None
            meta.chapter_count = sum(
                1 for p in document.paragraphs if _style_name(p).startswith("heading") and p.text.strip()
            ) or None
        except ConversionError:
            raise
        except Exception:
            pass
        return meta

    def extract(self, path: str, control: TaskControl | None = None) -> ExtractionResult:
        control = control or TaskControl()
        meta = BookMetadata.from_path(path)
        document = self._open(path)

        props = document.core_properties
        meta.title = props.title or None
        meta.author = props.author or None

        chapters: list[Chapter] = []
        current_title: str | None = None
        buffer: list[str] = []

        def flush():
            text = clean_text("\n".join(buffer))
            if text.strip() or current_title:
                chapters.append(Chapter(title=current_title, text=text))
            buffer.clear()

        for paragraph in document.paragraphs:
            control.pause_point()
            style = _style_name(paragraph)
            text = paragraph.text.strip()
            if style.startswith("heading") and text:
                flush()
                current_title = text
            elif text:
                buffer.append(paragraph.text)
        flush()
        control.progress(1.0)

        if not chapters:
            raise ConversionError("DOCX 中没有可提取的正文内容")
        meta.chapter_count = len(chapters)
        return ExtractionResult(metadata=meta, chapters=chapters)
